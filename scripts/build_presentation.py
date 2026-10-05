"""Build a seven-slide, five-minute Korean report from the published evidence."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
NAVY='14283F';BLUE='246BCE';GRAY='526276';BG='F6F8FC';WHITE='FFFFFF'
FONT='Noto Sans CJK KR'
notes=[]

def text(slide,x,y,w,h,content,size=20,color=NAVY,bold=False):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.03)
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
        p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(10)
        for r in p.runs:
            prop=r._r.get_or_add_rPr()
            ea=OxmlElement('a:ea');ea.set('typeface',FONT);prop.append(ea)
    return shape

def slide(title,kicker,duration,note):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    text(s,.55,.24,12.2,.35,kicker,12,BLUE,True)
    text(s,.55,.78,12.2,.75,title,29,bold=True)
    text(s,.58,7.09,11.4,.25,'unseen 지형 보행 학습  |  Isaac-Ant-v0  |  자체 평가 결과',10,GRAY)
    text(s,12.05,7.04,.7,.32,str(len(prs.slides)),12,GRAY)
    s.notes_slide.notes_text_frame.text=note
    notes.append((title,duration,note))
    return s

def picture(s,path,x,y,w,h):
    from PIL import Image
    p=ROOT/path
    with Image.open(p) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    return s.shapes.add_picture(str(p),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def table(s,x,y,w,headers,rows,widths=None,size=17):
    h=.55*(len(rows)+1)
    tbl=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
    if widths:
        for c,v in zip(tbl.columns,widths):c.width=Inches(w*v)
    for i,row in enumerate([headers]+rows):
        for j,value in enumerate(row):
            cell=tbl.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.10);cell.margin_top=Inches(.07)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(NAVY if i==0 else (WHITE if i%2 else 'EAF0F8'))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT;p.font.size=Pt(size);p.font.bold=i==0
                p.font.color.rgb=RGBColor.from_string(WHITE if i==0 else NAVY)
    return tbl

s=slide('처음 보는 지형에서도, 더 오래 전진할 수 있을까?','01  연구 문제',35,
'''과제의 목표는 특정 바닥에서 높은 점수를 얻는 것이 아니라 처음 보는 지형에서도 제한시간 동안 전진하는 Ant를 만드는 것입니다. 요철과 경사는 발이 닿는 높이를 바꾸고, 마찰은 같은 토크가 만드는 추진력을 바꿉니다. 우리는 다양한 지형을 경험시키고 주변 높이 정보를 제공한 뒤, 넘어짐과 초기 소환 문제를 구분해 보완했습니다. 지금 보여드리는 수치는 우리가 만든 맵과 팀 공통 맵의 자체 평가이며, 조교가 공개할 공식 unseen 결과는 아닙니다.''')
text(s,.6,1.9,5.3,2.8,'목표: 16초 안에 멀리 전진\n문제: 새 접촉면에서 균형·추진력 변화\n지표: 원본 보상 + 거리 + 생존율',23)
picture(s,'artifacts/media/self_eval/E3_SpawnLift.png',6.25,1.8,6.45,3.7)
text(s,.65,5.7,12,1,'로봇 형상·관절·USD는 유지하고, 지형 경험·관측·학습 보상을 설계했다.',21,BLUE,True)

s=slide('학습 맵과 자체 평가맵을 분리했다','02  실험 환경',45,
'''학습 지형은 요철, 파도, 경사, 높이가 다른 블록, 계단을 포함한 혼합의 다섯 구성입니다. 하나의 정책이 이 조건을 함께 경험합니다. 자체 평가 E1은 새로운 배치와 마찰에서 시작했고, E2에서는 연속 높낮이를 강조했습니다. E3에는 블록을 다시 넣어 블록 35%, 요철 25%, 파도 20%, 경사 20%로 만들었습니다. 지형 seed는 9317, 지면 마찰은 0.9와 0.75입니다. 학습과 다른 메시지만, 결과를 보고 개선에 사용했으므로 개발 검증 맵으로 분명히 구분합니다.''')
picture(s,'docs/ant_six_envs/assets/blocks_eval/preview.png',.55,1.7,7.0,3.65)
text(s,7.8,1.9,4.9,3.8,'E3 · 직접 구성한 평가맵\n블록 35% / 요철 25%\n파도 20% / 경사 상하 20%\nseed 9317 · 마찰 0.9 / 0.75\n8m 타일 · 전체 320×320m',20)
text(s,.65,5.6,12,1.1,'E1(9117) → E2(9217) → E3(9317)\n그림: 실제 생성 메시의 CPU 시각화, 세로 3배 확대',16,GRAY)

s=slide('정보를 늘리고, 실패의 비용을 명확히 했다','03  가설과 학습 방법',45,
'''첫 가설은 다양한 지형과 마찰을 경험하면 특정 바닥에 대한 의존을 줄일 수 있다는 것입니다. 두 번째는 몸 상태만이 아니라 앞의 높낮이를 알면 발 디딤에 활용할 수 있다는 것입니다. 그래서 60개 상태에 63개 지면 높이를 더했습니다. 마지막으로 토크 명령의 급격한 변화와 몸체 각속도에 작은 비용을 주고, 실제로 넘어질 때 사건당 마이너스 2를 주었습니다. 평가에서는 이 학습 전용 비용을 빼고 원본 일곱 보상만 사용합니다. 모든 변경의 개별 효과를 완전히 분리한 실험은 아닙니다.''')
table(s,.6,1.85,12.1,['가설','방법','고정·구분할 조건'],[
 ['다양한 접촉 경험','지형 혼합 + 폭·파도 수·마찰 변화','학습 seed와 평가 seed 분리'],
 ['앞쪽 높낮이 활용','60 상태 + 9×7 높이 = 123D','8개 effort 행동·로봇 형상 유지'],
 ['실패 감소가 전진에 도움','액션 변화·각속도·실패 비용','평가는 원본 Ant 7항 보상'],
],widths=[.25,.43,.32],size=18)
text(s,.75,4.7,11.8,1.25,'정책: PPO · actor/critic 400–200–100 · ELU\n센서: 카메라 영상이 아닌 이상적인 RayCaster 높이 질의',20)
text(s,.75,6.15,11.8,.55,'학습의 추가 비용: 액션 변화 −0.002 / roll·pitch 각속도 −0.01 / 실패 사건 −2',16,BLUE)

s=slide('통통 튐을 억제하기 전에, 실패 원인을 나눴다','04  관찰 → 진단 → 수정',45,
'''처음에는 통통 튀는 움직임이 문제라고 생각했습니다. 하지만 자세나 수직 움직임을 더 강하게 억제한 실험은 성능이 좋아지지 않았습니다. 이후 처음 소환될 때 발이 블록에 들어가는 장면을 관찰했습니다. 영 행동 진단으로 정책의 행동과 초기 접촉 문제를 분리했고, 시작 높이를 15센티미터 올렸을 때 겹침이 사라지고 착지하는 것을 확인했습니다. 이를 학습에도 반영했습니다. 단, 모든 지형에서 충돌을 보장하는 해결책은 아니며, 공정한 비교를 위해 두 모델 모두 같은 SpawnLift 평가를 사용했습니다.''')
table(s,.65,1.8,12,['검토한 접근','관찰한 결과','최종 판단'],[
 ['자세·수직 움직임 비용 강화','보상·생존율 개선이 확인되지 않음','최종 구성에서 제외'],
 ['초기 발·블록 겹침 진단','영 행동에서도 초기 반동·뒤집힘 관찰','정책 문제와 초기화 문제 구분'],
 ['초기 높이 +0.15m','겹침 감소와 착지 확인','학습에 적용; 동일 조건으로 비교'],
],widths=[.29,.43,.28],size=17)
text(s,.8,4.6,11.8,1.4,'핵심 아이디어\n자연스러워 보이는 동작보다, 누적 전진을 끊는 실패와 시작 조건을 먼저 확인한다.',23,BLUE,True)

s=slide('자체 평가: 같은 맵에서 생존과 전진이 함께 증가했다','05  E3-SpawnLift 결과',50,
'''두 모델 모두 E3-SpawnLift, seed 24, 100개 환경, 원본 보상으로 비교했습니다. 기존 Failure2의 평균 보상은 73.75, 최종 모델은 81.90입니다. 생존율은 82%에서 88%, 전진 거리는 78.66미터에서 82.79미터로 증가했습니다. 에피소드 길이도 늘었습니다. 이 결과는 실패를 줄여 더 오래 전진하는 방향과 일치하지만, 한 학습 seed의 관찰이므로 인과 효과나 일반적 우위를 확정하지 않습니다. 또 이 결과를 원래 reset 조건의 E3 점수와 합쳐서 개선량을 부풀리지 않았습니다.''')
table(s,.6,1.8,12.1,['지표','기존 Failure2','최종 Failure2Lift'],[
 ['원본 보상 mean ± std','73.75 ± 21.31','81.90 ± 17.05'],
 ['Episode steps mean ± std','890.92 ± 166.37','934.19 ± 95.38'],
 ['960-step 생존율','82%','88%'],
 ['+x 거리 mean ± std','78.66 ± 22.74m','82.79 ± 17.96m'],
],widths=[.40,.30,.30],size=18)
text(s,.75,5.1,11.8,.9,'두 모델 모두 seed 24 · 100/100 완료 · 같은 +0.15m reset\nstd: episode 간 모집단 표준편차, 학습 seed 간 변동이 아님',17,GRAY)
link=text(s,.75,6.1,11.8,.5,'실제 자체 평가 보행 영상 보기 ↗',20,BLUE,True)
link.click_action.hyperlink.address='https://github.com/Kim-KiSuk/Robotics-Simulation-Assignment-1/blob/main/artifacts/media/self_eval/E3_SpawnLift.mp4'

s=slide('팀 공통 자체 평가: 새로운 형상에서도 전진을 유지했다','06  고정 모델의 v2.1 평가',50,
'''최종 모델을 고정한 뒤 팀 공통 v2.1의 네 환경을 평가했습니다. Easy는 낮은 레일, Medium은 좁은 틈새, Hard는 반복 원기둥입니다. 보상은 순서대로 94.77, 91.24, 69.91, 88.64이고 생존율은 89%, 79%, 74%, 70%입니다. Hard의 평균 보상이 Medium보다 높아도 생존율은 낮습니다. 평균과 실패 빈도를 함께 봐야 합니다. Medium에서는 틈새로 ray가 지면을 만나지 않는 기록도 있었지만, 그 점이 실패의 주원인인지는 확인하지 못했습니다. 이 평가는 원본 world-Z 종료와 원본 reset을 사용하므로 앞의 E3와 점수를 직접 비교하지 않습니다.''')
table(s,.6,1.85,12.1,['환경','보상 mean ± std','생존율','평균 +x 거리'],[
 ['Seen Control','94.77 ± 32.68','89%','96.75m'],
 ['Unseen Easy','91.24 ± 40.22','79%','93.18m'],
 ['Unseen Medium','69.91 ± 35.04','74%','72.41m'],
 ['Unseen Hard','88.64 ± 45.20','70%','90.84m'],
],widths=[.28,.32,.16,.24],size=18)
text(s,.75,5.1,11.8,1.2,'모든 환경: seed 24 · 100/100 완료 · 동일 최종 checkpoint\n원본 보상·원본 종료·원본 reset / 관측·행동·센서는 학습 당시 유지',17,GRAY)
text(s,.75,6.35,11.8,.4,'팀 공통 자체 평가이며, 조교의 공식 unseen 평가 결과가 아니다.',16,BLUE,True)

s=slide('성능 수치와 함께, 무엇이 검증됐는지 남겼다','07  결론과 한계',30,
'''정리하면 다양한 접촉 경험과 높이 관측을 바탕으로, 작은 안정성 비용과 초기 소환 문제의 분리를 시도했습니다. 자체 공통 조건에서는 생존과 전진이 함께 증가했고, 새 지형에서도 전진은 유지됐습니다. 하지만 생존율 90% 목표에는 도달하지 못했고, 각 방법의 독립 효과나 여러 학습 seed의 재현성은 남은 한계입니다. 모든 환경 코드와 가중치, 평가 명령, 원자료와 영상을 저장소에 남겼습니다. 이번 발표는 높은 reward 순위보다 관찰에서 가설을 세우고 같은 조건에서 검증한 과정을 설명하는 데 초점을 맞췄습니다.''')
text(s,.8,1.95,11.7,2.7,'확인한 것\n• 동일 자체 맵에서 최종 모델의 생존·누적 전진 증가\n• 처음 보는 형상에서도 일정 수준의 전진 유지',24)
text(s,.8,4.35,11.7,1.8,'남은 한계\n• 단일 학습 seed · 개발 맵에 대한 반복 선택\n• 요소별 기여 미분리 · 생존율 90% 미달',22,GRAY)
text(s,.8,6.4,11.7,.4,'제출: GitHub 프로젝트 + 평가 명령어 TXT + 5분 PPT',18,BLUE,True)

OUT.mkdir(exist_ok=True)
prs.save(OUT/'Ant_Unseen_Terrain_5min.pptx')
elapsed=0;parts=['# 5분 발표 원고\n\n총 300초 계획. 실제 발표 전에 한 번 소리 내어 리허설한다. 조명·영상 재생 시간에 따라 조절한다.\n']
for i,(title,duration,note) in enumerate(notes,1):
    parts.append(f'## {i}. {title} ({elapsed//60}:{elapsed%60:02d}–{(elapsed+duration)//60}:{(elapsed+duration)%60:02d})\n\n{note}\n')
    elapsed+=duration
assert elapsed==300
(OUT/'SPEAKER_NOTES.md').write_text('\n'.join(parts))
print('Saved 7 slides; 300-second speaker plan.')
